"""Füllt die Sponsorenkarte der Schule (Word-Vorlage) und exportiert sie optional als PDF."""
import copy
import io
import shutil
import subprocess
import tempfile
from pathlib import Path

from docx import Document
from docx.shared import Pt, RGBColor

from .logic import erloes_cent
from .settings import as_int
from .util import fmt_euro

TEMPLATE = Path(__file__).parent / "resources" / "sponsorenkarte.docx"
OPT_IN = "Opt-In per Email-Bestätigung"
MIN_ROWS = 5


def _set(paragraph, text):
    """Text eines Absatzes setzen und dabei die Formatierung des ersten Laufs behalten."""
    runs = paragraph.runs
    if runs:
        runs[0].text = text
        for r in runs[1:]:
            r.text = ""
    else:
        paragraph.add_run(text)


def _plain(paragraph, size=10):
    """Datenzeilen einheitlich: normal, schwarz, 10 pt (die Vorlage ist dort teils fett/farbig)."""
    for run in paragraph.runs:
        run.font.size = Pt(size)
        run.font.bold = False
        run.font.color.rgb = RGBColor(0, 0, 0)


def _cell_lines(cell, lines):
    paras = cell.paragraphs
    for i, line in enumerate(lines):
        if i < len(paras):
            _set(paras[i], line)
        else:
            cell.add_paragraph(line)
        _plain(cell.paragraphs[i])


def build_docx(sponsors, s, result_known):
    doc = Document(str(TEMPLATE))
    head, table = doc.tables[0], doc.tables[1]
    runden = as_int(s, "runden")

    left = head.rows[0].cells[0].paragraphs
    _set(left[0], f"Klasse: {s.get('klasse', '')}")
    _set(left[2], f"Sponsorenkarte von: {s.get('karten_name', '')}")
    _set(head.rows[0].cells[1].paragraphs[0],
         f"Anzahl der gelaufenen Runden: {runden if result_known else '____'}")

    _set(table.rows[0].cells[0].paragraphs[1], f"({as_int(s, 'rundenlaenge_m', 100)}m)")

    # Datenzeilen bei Bedarf ergänzen (Kopie der ersten Datenzeile)
    template_tr = table.rows[1]._tr
    while len(table.rows) - 1 < max(MIN_ROWS, len(sponsors)):
        template_tr.getparent().append(copy.deepcopy(template_tr))

    for idx, sp in enumerate(sponsors, start=1):
        cells = table.rows[idx].cells
        betrag = fmt_euro(sp["betrag_cent"])
        if sp["hoechst_cent"]:
            betrag += f" (max. {fmt_euro(sp['hoechst_cent'])})"
        _cell_lines(cells[0], [betrag])
        _cell_lines(cells[1], [sp["name"]])
        _cell_lines(cells[2], [OPT_IN if sp["bestaetigt_durch"] == "link" else ""])
        _cell_lines(cells[3], [fmt_euro(erloes_cent(sp, runden)) if result_known else ""])
        if sp["quittung"]:
            _cell_lines(cells[4], [sp["quittung_name"] or sp["name"], f"{sp['strasse'] or ''}, {sp['plz_ort'] or ''}".strip(", ")])
        else:
            _cell_lines(cells[4], [""])

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf


def docx_to_pdf(docx_bytes, soffice_bin="soffice"):
    """Konvertiert per LibreOffice (headless). Gibt None zurück, wenn LibreOffice fehlt."""
    if not shutil.which(soffice_bin) and not Path(soffice_bin).exists():
        return None
    with tempfile.TemporaryDirectory() as tmp:
        src = Path(tmp) / "sponsorenkarte.docx"
        src.write_bytes(docx_bytes)
        subprocess.run(
            [soffice_bin, "--headless", "--norestore", f"-env:UserInstallation=file://{tmp}/lo",
             "--convert-to", "pdf", "--outdir", tmp, str(src)],
            check=True, timeout=90, capture_output=True,
        )
        pdf = Path(tmp) / "sponsorenkarte.pdf"
        return pdf.read_bytes() if pdf.exists() else None
